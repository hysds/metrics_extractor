#!/usr/bin/env python
"""
Build a combined Excel workbook with all PGE three-level breakdowns for a CRID.

The workbook has:
- Summary sheet: per-PGE totals, weighted avg container/job runtime, bar chart
- "All PGEs (unified)" sheet: the aggregated CSV with normalized level_1/2/3 columns
- One sheet per PGE type (L0B, RSLC, GSLC, GCOV, INSAR, L3_SM) with native schema

Usage:
    build_crid_workbook.py --input_dir ~/dev/nisar/tmp \\
        --crid X05013 --cluster POP1 --version r05.01.3 \\
        [--output OUT.xlsx]
"""

import argparse
import csv
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)

FNAME_RE = re.compile(r"job_three_level_breakdown_job_SCIFLO_([A-Z0-9_]+?)_(release|pcm).*\.csv$")


def discover_csvs(input_dir: Path, version_filter: str | None):
    matches = {}
    for p in input_dir.glob("job_three_level_breakdown_*.csv"):
        m = FNAME_RE.search(p.name)
        if not m:
            continue
        pge = m.group(1).rstrip("_")
        if version_filter and version_filter.replace(".", "") not in p.name.replace(".", ""):
            continue
        if pge not in matches or p.stat().st_mtime > matches[pge].stat().st_mtime:
            matches[pge] = p
    return matches


def find_aggregated(input_dir: Path, crid: str, cluster: str):
    """Locate the aggregated CSV produced by combine_breakdowns.py, or None."""
    candidates = sorted(input_dir.glob(f"all_pge_three_level_breakdown_{crid}_{cluster}_*.csv"), reverse=True)
    return candidates[0] if candidates else None


def style_header_row(ws, row_num, num_cols):
    for i in range(1, num_cols + 1):
        c = ws.cell(row=row_num, column=i)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL


def autofit(ws, default=14):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = default
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
            v = row[0]
            if v is not None:
                max_len = max(max_len, min(len(str(v)) + 2, 45))
        ws.column_dimensions[letter].width = max_len


def write_csv_to_sheet(ws, csv_path, numeric_cols=None):
    numeric_cols = numeric_cols or set()
    with open(csv_path) as f:
        rows = list(csv.reader(f))
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if r_idx == 1:
                cell.value = val
            elif c_idx in numeric_cols and val:
                try:
                    cell.value = float(val)
                except ValueError:
                    cell.value = val
            else:
                cell.value = val
    if rows:
        style_header_row(ws, 1, len(rows[0]))
    autofit(ws)


def numeric_cols_for_header(header):
    """Guess which 1-indexed columns are numeric from header names."""
    cols = set()
    for i, h in enumerate(header, start=1):
        if any(k in h for k in ["_m", "_gb", "_mbps", "count", "_avg", "days", "rcid"]):
            cols.add(i)
    return cols


def build_summary(wb, csvs, crid, cluster, version, days_back):
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "NISAR PGE Three-Level Breakdown Metrics"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"CRID: {crid}" + (f" ({version})" if version else "")
    ws["A3"] = f"Cluster: {cluster}"
    ws["A4"] = f"Lookback window: {days_back} days" if days_back else ""
    ws["A5"] = f"Report Date: {datetime.now().strftime('%Y-%m-%d')}"

    ws["A7"] = "Per-PGE Summary"
    ws["A7"].font = BOLD
    hdr_row = 8
    hdrs = ["PGE Type", "Breakdown Groups", "Total Jobs", "Weighted Avg Container (min)", "Weighted Avg Job (min)"]
    for i, h in enumerate(hdrs, start=1):
        ws.cell(row=hdr_row, column=i).value = h
    style_header_row(ws, hdr_row, len(hdrs))

    row = hdr_row + 1
    pge_summary = []
    for pge, path in sorted(csvs.items()):
        with open(path) as f:
            reader = csv.DictReader(f)
            total_count = 0
            weighted_c = 0.0
            weighted_j = 0.0
            group_count = 0
            for r in reader:
                c = int(r.get("count", 0))
                group_count += 1
                total_count += c
                weighted_c += float(r.get("container_runtime_m", 0)) * c
                weighted_j += float(r.get("job_runtime_m", 0)) * c
        avg_c = weighted_c / total_count if total_count else 0
        avg_j = weighted_j / total_count if total_count else 0
        ws.cell(row=row, column=1).value = pge
        ws.cell(row=row, column=2).value = group_count
        ws.cell(row=row, column=3).value = total_count
        ws.cell(row=row, column=4).value = round(avg_c, 2)
        ws.cell(row=row, column=5).value = round(avg_j, 2)
        pge_summary.append((pge, group_count, total_count, avg_c, avg_j))
        row += 1

    ws.cell(row=row, column=1).value = "TOTAL"
    ws.cell(row=row, column=1).font = BOLD
    ws.cell(row=row, column=2).value = sum(p[1] for p in pge_summary)
    ws.cell(row=row, column=3).value = sum(p[2] for p in pge_summary)
    ws.cell(row=row, column=2).font = BOLD
    ws.cell(row=row, column=3).font = BOLD

    ws.cell(row=row + 2, column=1).value = "Notes"
    ws.cell(row=row + 2, column=1).font = BOLD
    notes = [
        "• RSLC/GSLC/GCOV/INSAR/L3_SM use three-level breakdown: beam_name → coverage → acquisition_mode",
        "• L0B uses three-level breakdown: rcid → beam_mode → diagnostic_mode (runs pre-focus)",
        "• 'Container' = PGE container wall time (min); 'Job' = total HySDS job duration (min, includes stage in/out)",
        "• Stage I/O columns populated for RSLC/GSLC/GCOV/INSAR/L3_SM; empty for L0B",
        "• 'All PGEs (unified)' sheet aggregates all groups with normalized level_1/level_2/level_3 columns",
    ]
    for i, n in enumerate(notes, start=row + 3):
        ws.cell(row=i, column=1).value = n

    autofit(ws)
    for col in "ABCDE":
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 14, 22)

    if pge_summary:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Avg Container Runtime by PGE Type"
        chart.y_axis.title = "Avg Container (min)"
        chart.x_axis.title = "PGE Type"
        data = Reference(ws, min_col=4, min_row=hdr_row, max_col=4, max_row=hdr_row + len(pge_summary))
        cats = Reference(ws, min_col=1, min_row=hdr_row + 1, max_row=hdr_row + len(pge_summary))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 18
        chart.height = 10
        ws.add_chart(chart, "G7")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input_dir", required=True, help="Directory with per-PGE breakdown CSVs")
    ap.add_argument("--crid", required=True, help="CRID (e.g., X05013)")
    ap.add_argument("--cluster", required=True, help="Cluster name (e.g., POP1)")
    ap.add_argument("--version", help="Release version filter (e.g., r05.01.3)")
    ap.add_argument("--days_back", type=int, help="Lookback window in days — for display only")
    ap.add_argument("--output", help="Output xlsx path")
    args = ap.parse_args()

    input_dir = Path(args.input_dir).expanduser()
    if not input_dir.is_dir():
        print(f"ERROR: {input_dir} is not a directory", file=sys.stderr)
        sys.exit(1)

    csvs = discover_csvs(input_dir, args.version)
    if not csvs:
        print(f"ERROR: no matching breakdown CSVs found in {input_dir}", file=sys.stderr)
        sys.exit(1)

    agg_csv = find_aggregated(input_dir, args.crid, args.cluster)
    if not agg_csv:
        print(f"WARNING: no aggregated CSV found (all_pge_three_level_breakdown_{args.crid}_{args.cluster}_*.csv). Run combine_breakdowns.py first for the unified sheet.", file=sys.stderr)

    out_path = Path(args.output) if args.output else input_dir / f"ALL_PGE_three_level_breakdown_{args.crid}_{args.cluster}_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

    wb = openpyxl.Workbook()

    build_summary(wb, csvs, args.crid, args.cluster, args.version, args.days_back)

    if agg_csv:
        ws_all = wb.create_sheet("All PGEs (unified)")
        with open(agg_csv) as f:
            header = next(csv.reader(f))
        write_csv_to_sheet(ws_all, agg_csv, numeric_cols=numeric_cols_for_header(header))

    for pge, path in sorted(csvs.items()):
        ws_pge = wb.create_sheet(pge)
        with open(path) as f:
            header = next(csv.reader(f))
        write_csv_to_sheet(ws_pge, path, numeric_cols=numeric_cols_for_header(header))

    wb.save(out_path)
    print(f"Saved: {out_path}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
