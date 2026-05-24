#!/usr/bin/env python
"""
PGE Version Comparison

Compares execution times between two PGE versions using CSV files from
pge_execution_time_extractor.py. Produces an Excel report with matched
job comparisons.

Jobs are matched by their match_id (product timestamps).

Usage:
    compare_pge_versions.py --old_csv=v1.csv --new_csv=v2.csv --output=comparison.xlsx
"""

import csv
import re
import sys
from argparse import ArgumentParser

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side
    from openpyxl.chart import BarChart, Reference
except ImportError:
    import subprocess
    subprocess.check_call(['pip', 'install', 'openpyxl', '-q'])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side
    from openpyxl.chart import BarChart, Reference


def extract_version(job_id):
    """Extract version string from job_id."""
    match = re.search(r'(release-r[\d.]+|pcm_r[\d.]+_pge_r[\d.]+)', job_id)
    return match.group(1) if match else "unknown"


def read_csv(filepath):
    """Read CSV and return jobs dict keyed by match_id."""
    jobs = {}
    version = None
    pge_type = None
    data_day = None

    with open(filepath, 'r') as f:
        reader = csv.reader(f)
        in_details = False
        header = None

        for row in reader:
            if not row:
                continue

            # Parse summary section for metadata
            if row[0] == "# Summary Statistics":
                continue

            if row[0] == "# Individual Job Details":
                in_details = True
                continue

            # Get header row
            if in_details and row[0] == "job_id":
                header = row
                continue

            # Parse summary row for pge_type and data_day
            if not in_details and row[0] in ['L0B', 'RSLC', 'GSLC', 'GCOV', 'INSAR', 'L3_SM']:
                pge_type = row[0]
                data_day = row[2] if len(row) > 2 else None

            # Parse job details
            if in_details and header and row[0] != "job_id":
                job_id = row[0]
                match_id = row[1]
                instance_type = row[3]
                pge_time = float(row[4])
                pcm_time = float(row[5])

                if version is None:
                    version = extract_version(job_id)

                jobs[match_id] = {
                    'job_id': job_id,
                    'instance_type': instance_type,
                    'pge_time': pge_time,
                    'pcm_time': pcm_time,
                }

    return jobs, version, pge_type, data_day


def create_excel(old_jobs, new_jobs, old_ver, new_ver, pge_type, data_day, output):
    """Create comparison Excel workbook."""
    # Find matching jobs with same instance type
    matches = []
    for match_id in set(old_jobs.keys()) & set(new_jobs.keys()):
        old = old_jobs[match_id]
        new = new_jobs[match_id]
        if old['instance_type'] != new['instance_type']:
            continue

        diff = new['pge_time'] - old['pge_time']
        pct = (diff / old['pge_time'] * 100) if old['pge_time'] else 0

        matches.append({
            'match_id': match_id,
            'instance_type': old['instance_type'],
            'old_time': old['pge_time'],
            'new_time': new['pge_time'],
            'diff': diff,
            'pct': pct,
            'old_job_id': old['job_id'],
            'new_job_id': new['job_id'],
        })

    if not matches:
        print("Error: No matching jobs found")
        print(f"Old CSV match_ids (sample): {list(old_jobs.keys())[:3]}")
        print(f"New CSV match_ids (sample): {list(new_jobs.keys())[:3]}")
        sys.exit(1)

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"

    ws['A1'] = f"{pge_type or 'PGE'} Execution Time Comparison"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"Data Day: {data_day or 'N/A'}"
    ws['A3'] = f"Comparing: {old_ver} vs {new_ver}"

    faster = sum(1 for m in matches if m['diff'] < 0)
    slower = sum(1 for m in matches if m['diff'] > 0)
    avg_diff = sum(m['diff'] for m in matches) / len(matches)
    avg_pct = sum(m['pct'] for m in matches) / len(matches)

    ws['A5'] = "Overall Summary"
    ws['A5'].font = Font(bold=True)
    ws['A6'] = "Total matched jobs:"
    ws['B6'] = len(matches)
    ws['A7'] = f"Jobs where {new_ver} is faster:"
    ws['B7'] = f"{faster} ({faster/len(matches)*100:.1f}%)"
    ws['A8'] = f"Jobs where {new_ver} is slower:"
    ws['B8'] = f"{slower} ({slower/len(matches)*100:.1f}%)"
    ws['A9'] = "Average difference:"
    ws['B9'] = f"{avg_diff:+.2f} min"
    ws['A10'] = "Average % change:"
    ws['B10'] = f"{avg_pct:+.1f}%"

    # Summary by instance type
    by_inst = {}
    for m in matches:
        inst = m['instance_type']
        if inst not in by_inst:
            by_inst[inst] = []
        by_inst[inst].append(m)

    ws['A12'] = "By Instance Type"
    ws['A12'].font = Font(bold=True)

    headers = ["Instance Type", "Count", f"Avg {old_ver}", f"Avg {new_ver}", "Avg Diff", "Avg %"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=13, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    row = 14
    for inst in sorted(by_inst.keys()):
        jobs = by_inst[inst]
        avg_old = sum(j['old_time'] for j in jobs) / len(jobs)
        avg_new = sum(j['new_time'] for j in jobs) / len(jobs)
        avg_d = sum(j['diff'] for j in jobs) / len(jobs)
        avg_p = sum(j['pct'] for j in jobs) / len(jobs)

        ws.cell(row=row, column=1, value=inst).border = border
        ws.cell(row=row, column=2, value=len(jobs)).border = border
        ws.cell(row=row, column=3, value=round(avg_old, 2)).border = border
        ws.cell(row=row, column=4, value=round(avg_new, 2)).border = border

        diff_cell = ws.cell(row=row, column=5, value=round(avg_d, 2))
        diff_cell.border = border
        diff_cell.fill = green_fill if avg_d < 0 else red_fill

        pct_cell = ws.cell(row=row, column=6, value=f"{avg_p:+.1f}%")
        pct_cell.border = border
        pct_cell.fill = green_fill if avg_p < 0 else red_fill

        row += 1

    for col, width in [('A', 35), ('B', 20), ('C', 18), ('D', 18), ('E', 12), ('F', 12)]:
        ws.column_dimensions[col].width = width

    # Sheet 2: Matched Comparison
    ws2 = wb.create_sheet("Matched Comparison")

    headers = ["Match ID", "Instance Type", f"{old_ver} (min)", f"{new_ver} (min)", "Diff (min)", "% Change"]
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    for i, m in enumerate(sorted(matches, key=lambda x: x['match_id']), 2):
        ws2.cell(row=i, column=1, value=m['match_id']).border = border
        ws2.cell(row=i, column=2, value=m['instance_type']).border = border
        ws2.cell(row=i, column=3, value=round(m['old_time'], 2)).border = border
        ws2.cell(row=i, column=4, value=round(m['new_time'], 2)).border = border

        diff_cell = ws2.cell(row=i, column=5, value=round(m['diff'], 2))
        diff_cell.border = border
        diff_cell.fill = green_fill if m['diff'] < 0 else red_fill

        pct_cell = ws2.cell(row=i, column=6, value=round(m['pct'], 1))
        pct_cell.border = border
        pct_cell.fill = green_fill if m['pct'] < 0 else red_fill

    for col, width in [('A', 45), ('B', 14), ('C', 18), ('D', 18), ('E', 12), ('F', 12)]:
        ws2.column_dimensions[col].width = width

    # Sheet 3: Raw Data - Old
    ws3 = wb.create_sheet("Raw Data - Old")
    headers = ["Job ID", "Match ID", "Instance Type", "PGE Time (min)", "PCM Time (min)"]
    for col, h in enumerate(headers, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for i, (mid, job) in enumerate(sorted(old_jobs.items()), 2):
        ws3.cell(row=i, column=1, value=job['job_id'])
        ws3.cell(row=i, column=2, value=mid)
        ws3.cell(row=i, column=3, value=job['instance_type'])
        ws3.cell(row=i, column=4, value=round(job['pge_time'], 2))
        ws3.cell(row=i, column=5, value=round(job['pcm_time'], 2))

    ws3.column_dimensions['A'].width = 100
    ws3.column_dimensions['B'].width = 45

    # Sheet 4: Raw Data - New
    ws4 = wb.create_sheet("Raw Data - New")
    for col, h in enumerate(headers, 1):
        cell = ws4.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for i, (mid, job) in enumerate(sorted(new_jobs.items()), 2):
        ws4.cell(row=i, column=1, value=job['job_id'])
        ws4.cell(row=i, column=2, value=mid)
        ws4.cell(row=i, column=3, value=job['instance_type'])
        ws4.cell(row=i, column=4, value=round(job['pge_time'], 2))
        ws4.cell(row=i, column=5, value=round(job['pcm_time'], 2))

    ws4.column_dimensions['A'].width = 100
    ws4.column_dimensions['B'].width = 45

    # Sheet 5: Chart
    ws5 = wb.create_sheet("Chart")
    ws5['A1'] = "Instance Type"
    ws5['B1'] = f"{old_ver} Avg"
    ws5['C1'] = f"{new_ver} Avg"

    for col in range(1, 4):
        ws5.cell(row=1, column=col).font = header_font
        ws5.cell(row=1, column=col).fill = header_fill

    row = 2
    for inst in sorted(by_inst.keys()):
        jobs = by_inst[inst]
        ws5.cell(row=row, column=1, value=inst)
        ws5.cell(row=row, column=2, value=round(sum(j['old_time'] for j in jobs) / len(jobs), 2))
        ws5.cell(row=row, column=3, value=round(sum(j['new_time'] for j in jobs) / len(jobs), 2))
        row += 1

    chart = BarChart()
    chart.type = "col"
    chart.title = "Average Execution Time by Instance Type"
    chart.y_axis.title = "Minutes"
    chart.add_data(Reference(ws5, min_col=2, min_row=1, max_row=row-1, max_col=3), titles_from_data=True)
    chart.set_categories(Reference(ws5, min_col=1, min_row=2, max_row=row-1))
    chart.width = 15
    chart.height = 10
    ws5.add_chart(chart, "E2")

    wb.save(output)

    print(f"Saved: {output}")
    print(f"\nSummary:")
    print(f"  Matched jobs: {len(matches)}")
    print(f"  {new_ver} faster: {faster} ({faster/len(matches)*100:.1f}%)")
    print(f"  {new_ver} slower: {slower} ({slower/len(matches)*100:.1f}%)")
    print(f"  Average change: {avg_diff:+.2f} min ({avg_pct:+.1f}%)")


def main():
    parser = ArgumentParser(description="Compare PGE execution times between versions")
    parser.add_argument("--old_csv", required=True, help="CSV for baseline version")
    parser.add_argument("--new_csv", required=True, help="CSV for new version")
    parser.add_argument("--output", default="pge_comparison.xlsx", help="Output Excel file")
    parser.add_argument("--old_version", help="Label for old version")
    parser.add_argument("--new_version", help="Label for new version")
    args = parser.parse_args()

    print(f"Reading: {args.old_csv}")
    old_jobs, old_ver, pge_type_old, data_day_old = read_csv(args.old_csv)
    print(f"  {len(old_jobs)} jobs, version: {old_ver}")

    print(f"Reading: {args.new_csv}")
    new_jobs, new_ver, pge_type_new, data_day_new = read_csv(args.new_csv)
    print(f"  {len(new_jobs)} jobs, version: {new_ver}")

    create_excel(
        old_jobs, new_jobs,
        args.old_version or old_ver,
        args.new_version or new_ver,
        pge_type_old or pge_type_new,
        data_day_old or data_day_new,
        args.output
    )


if __name__ == "__main__":
    main()
